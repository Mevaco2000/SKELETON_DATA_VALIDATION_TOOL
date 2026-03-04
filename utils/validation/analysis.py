"""Analysis of image groups and visualization."""

import os
import shutil
import json
import random
from typing import List, Dict, Any, Tuple
import numpy as np
import cv2
import matplotlib.pyplot as plt
from tqdm import tqdm
from openpyxl import Workbook

from .image_validation import YPImageValidation


def generate_group_visualizations(
    groups: List[List[int]],
    filenames: List[str],
    image_folder: str,
    label_folder: str,
    output_folder: str
) -> None:
    """Generate visualizations for image groups.
    
    Args:
        groups: List of groups, each group is a list of image indices
        filenames: List of image filenames
        image_folder: Path to folder with images
        label_folder: Path to folder with labels
        output_folder: Path to save visualizations
    """
    if os.path.exists(output_folder):
        shutil.rmtree(output_folder)

    os.makedirs(output_folder)

    for i, group in enumerate(tqdm(groups, desc="Generating visualizations")):

        base_idx = group[0]
        base_img = cv2.imread(os.path.join(image_folder, filenames[base_idx]))
        image = cv2.cvtColor(base_img, cv2.COLOR_BGR2RGB)

        legend_entries = []

        for idx in group:
            img_name = filenames[idx]
            label_name = os.path.splitext(img_name)[0] + ".txt"
            label_path = os.path.join(label_folder, label_name)

            if not os.path.exists(label_path):
                continue

            kp = YPImageValidation.load_keypoints(label_path)
            color = tuple(random.randint(0, 255) for _ in range(3))

            image = YPImageValidation.draw_keypoints_on_image(image, kp, color)
            legend_entries.append((f"(idx={idx})", color))

        output_path = os.path.join(output_folder, f"group_{i}.png")

        plt.figure(figsize=(8, 8))
        plt.imshow(image)
        plt.axis("off")

        for name, color in legend_entries:
            plt.plot([], [], marker='o',
                     color=np.array(color)/255,
                     linestyle='None',
                     label=name)

        plt.legend(fontsize=6, loc="upper right")
        plt.tight_layout()
        plt.savefig(output_path)
        plt.close()


def save_groups_analysis(
    groups: List[List[int]],
    filenames: List[str],
    label_folder: str,
    output_json: str
) -> None:
    """Save group analysis to JSON file.
    
    Args:
        groups: List of groups
        filenames: List of image filenames
        label_folder: Path to folder with labels
        output_json: Path to save JSON analysis
    """
    data = {}

    for i, group in enumerate(tqdm(groups, desc="Saving group analysis")):
        group_key = f"group_{i}"
        data[group_key] = []

        for idx in sorted(group):

            img_name = filenames[idx]
            label_path = os.path.join(
                label_folder,
                os.path.splitext(img_name)[0] + ".txt"
            )

            entry = {"filename": img_name, "index": int(idx)}

            if os.path.exists(label_path):
                kp = YPImageValidation.load_keypoints(label_path)
                dist_vector = YPImageValidation.compute_sequential_distances(kp)
                entry["distance_vector"] = [
                    round(float(d), 2) for d in dist_vector
                ]

            data[group_key].append(entry)

    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4)


def analyze_hidden_keypoints(label_folder: str, output_json: str) -> None:
    """Analyze hidden keypoints in labels.
    
    Args:
        label_folder: Path to folder with labels
        output_json: Path to save analysis
    """
    results = []

    label_files = [
        f for f in os.listdir(label_folder)
        if f.endswith(".txt")
    ]

    for filename in tqdm(label_files, desc="Analyzing hidden keypoints"):

        path = os.path.join(label_folder, filename)

        with open(path, "r") as f:
            lines = f.readlines()

        hidden_keypoints = []

        for object_id, line in enumerate(lines):
            parts = line.strip().split()
            keypoints_data = list(map(float, parts[5:]))

            for i in range(0, len(keypoints_data), 3):
                if int(keypoints_data[i + 2]) == 1:
                    hidden_keypoints.append({
                        "object_id": object_id,
                        "keypoint_index": i // 3
                    })

        if hidden_keypoints:
            results.append({
                "filename": filename,
                "hidden_keypoints": hidden_keypoints
            })

    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=4)


def export_groups_analysis_to_excel(json_path: str, output_path: str) -> None:
    """Export group analysis to Excel with statistics.
    
    For each group:
    - Calculates mean for each distance (d1, d2, ...)
    - Calculates global group mean
    - Sorts groups in descending order by global mean
    - Saves to Excel
    
    Args:
        json_path: Path to groups_analysis.json
        output_path: Path for output Excel file
    """
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    processed_groups = []

    for group_name, group_items in data.items():

        vectors = [
            item.get("distance_vector", [])
            for item in group_items
            if "distance_vector" in item
        ]

        if not vectors:
            continue

        max_len = max(len(v) for v in vectors)

        # Pad vectors to same length
        padded = []
        for v in vectors:
            padded_vec = v + [np.nan] * (max_len - len(v))
            padded.append(padded_vec)

        padded = np.array(padded, dtype=float)

        # Mean for each keypoint distance
        mean_per_distance = np.nanmean(padded, axis=0)

        # Global group mean
        global_mean = float(np.nanmean(padded))

        processed_groups.append({
            "group_name": group_name,
            "items": group_items,
            "mean_per_distance": mean_per_distance,
            "global_mean": global_mean
        })

    # Sort descending by global mean
    processed_groups.sort(
        key=lambda x: x["global_mean"],
        reverse=True
    )

    # Export to Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Groups_sorted_by_mean"

    current_row = 1

    for group in tqdm(processed_groups, desc="Exporting sorted groups"):

        ws.cell(row=current_row, column=1,
                value=f"{group['group_name']} (mean={round(group['global_mean'], 4)})")
        current_row += 1

        ws.cell(row=current_row, column=1, value="index")

        max_vector_len = len(group["mean_per_distance"])

        for i in range(max_vector_len):
            ws.cell(row=current_row, column=2 + i, value=f"d{i+1}")

        current_row += 1

        # Data rows
        for item in group["items"]:

            ws.cell(row=current_row, column=1, value=item.get("index"))

            vector = item.get("distance_vector", [])
            for i in range(len(vector)):
                ws.cell(row=current_row, column=2 + i, value=vector[i])

            current_row += 1

        # Mean row
        ws.cell(row=current_row, column=1, value="MEAN_PER_DISTANCE")
        for i, val in enumerate(group["mean_per_distance"]):
            ws.cell(row=current_row, column=2 + i,
                    value=float(round(val, 4)))

        current_row += 2

    wb.save(output_path)


# ---------------------------------------------------------------------------
# group visualization and analysis functions
# ---------------------------------------------------------------------------

